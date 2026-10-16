#!/usr/bin/env python3
"""Prepara o relatório exportado do OPA para publicar no painel.

- Lê CSV (UTF-8 ou Windows-1252; vírgula, ponto e vírgula ou tab) ou XLSX.
- Mantém TODAS as colunas do relatório, inclusive Id do Cliente, Bairro e Condomínio (quando existirem),
  exceto as que identificam a pessoa (contato, número do contato, nome do cliente, telefone, CPF, e-mail,
  endereço etc.). Colunas novas ficam como colunas extras.
- Normaliza datas para AAAA-MM-DD HH:MM.
- Com --acrescentar, junta ao arquivo de saída existente sem duplicar (chave: Protocolo ou id).

Exemplos:
  python tools/preparar_dados.py relatorio.xlsx
  python tools/preparar_dados.py relatorio.csv --acrescentar
  python tools/preparar_dados.py relatorio.csv -o data/atendimentos.csv

XLSX exige:  pip install openpyxl   (CSV não precisa de nada além do Python 3.8+)
"""
import argparse, csv, io, re, sys, unicodedata
from datetime import datetime, date

ALIASES = {
    "id": ["id"], "protocolo": ["protocolo"], "status": ["status"], "canal": ["canal"],
    "atendente": ["ultimo atendente", "atendente"], "departamento": ["ultimo departamento", "departamento"],
    "motivos": ["motivos do atendimento", "motivos", "motivo"],
    "abertoEm": ["aberto em"], "assumidoEm": ["assumido em"], "encerradoEm": ["encerrado em"],
    "avaliacao": ["avaliacao", "avaliacao do cliente", "nota"],
    "tempoAtendimento": ["tempo em atendimento"], "tempoEspera": ["tempo em espera"],
    "etiquetas": ["etiquetas do atendimento", "etiquetas"], "abertoPor": ["aberto por"], "encerradoPor": ["encerrado por"],
    "mensagensEnviadas": ["mensagens enviadas"], "mensagensRecebidas": ["mensagens recebidas"],
    "aiEvaluation": ["aievaluation", "ai evaluation"],
    "waitingTimeInLastDepartment": ["waitingtimeinlastdepartment", "waiting time in last department"],
    "idCliente": ["id do cliente", "id cliente", "codigo do cliente", "cod cliente", "codigo cliente"],
    "bairro": ["bairro", "bairro do cliente"], "condominio": ["condominio", "condominio do cliente"],
    "cidade": ["cidade", "cidade do cliente", "municipio"], "regiao": ["regiao", "regiao do cliente", "localidade"],
}
GEO_HINTS = [("bairro", "bairro"), ("condominio", "condominio"), ("cidade", "cidade")]
# Termos que, num cabeçalho NÃO reconhecido, indicam dado do cliente (a coluna é descartada).
SENSIVEIS = ["contato", "cliente", "telefone", "celular", "fone", "cpf", "cnpj", "email", "e mail",
             "endereco", "documento", "nome", "cep", "whatsapp", "rg"]
HEADERS = {
    "id": "id", "protocolo": "Protocolo", "status": "Status", "canal": "Canal", "atendente": "Último atendente",
    "departamento": "Último departamento", "motivos": "Motivos do atendimento", "abertoEm": "Aberto em",
    "assumidoEm": "Assumido em", "encerradoEm": "Encerrado em", "avaliacao": "Avaliação",
    "tempoAtendimento": "Tempo em atendimento", "tempoEspera": "Tempo em espera",
    "etiquetas": "Etiquetas do atendimento", "abertoPor": "Aberto por", "encerradoPor": "Encerrado por",
    "mensagensEnviadas": "Mensagens enviadas", "mensagensRecebidas": "Mensagens recebidas",
    "aiEvaluation": "aiEvaluation", "waitingTimeInLastDepartment": "waitingTimeInLastDepartment",
    "idCliente": "Id do Cliente", "bairro": "Bairro", "condominio": "Condomínio", "cidade": "Cidade", "regiao": "Região",
}
FIELDS = list(HEADERS)
DATE_FIELDS = ("abertoEm", "assumidoEm", "encerradoEm")


def norm(s):
    s = unicodedata.normalize("NFD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def to_iso(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else v.strftime("%Y-%m-%d 00:00")
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:[ T,]+(\d{1,2}):(\d{2}))?", s)
    if m:
        d, mo, y, h, mi = m.groups()
        return f"{int(y):04d}-{int(mo):02d}-{int(d):02d} {int(h or 0):02d}:{int(mi or 0):02d}"
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]+(\d{1,2}):(\d{2}))?", s)
    if m:
        y, mo, d, h, mi = m.groups()
        return f"{int(y):04d}-{int(mo):02d}-{int(d):02d} {int(h or 0):02d}:{int(mi or 0):02d}"
    return s


def sensivel(h):
    k = " " + norm(h) + " "
    return any(" " + t + " " in k for t in SENSIVEIS)


def read_table(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Para ler .xlsx instale o openpyxl:  pip install openpyxl   (ou exporte como CSV)")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = next((wb[n] for n in wb.sheetnames if n.lower() == "dados"), wb[wb.sheetnames[0]])
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = open(path, "rb").read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
    head = text.splitlines()[0] if text else ""
    delim = max([",", ";", "\t"], key=head.count)
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def canonicalize(table):
    for hi, row in enumerate(table[:10]):
        m = {}
        for ci, h in enumerate(row):
            k = norm(h if h is not None else "")
            for f, al in ALIASES.items():
                if f not in m and k in al:
                    m[f] = ci
                    break
            else:
                for f, word in GEO_HINTS:  # ex.: "Endereço - Bairro"
                    if f not in m and (" " + word + " ") in (" " + k + " "):
                        m[f] = ci
                        break
        if len(m) >= 3 and "canal" in m:
            break
    else:
        sys.exit("Não encontrei o cabeçalho (Canal, Último atendente, Assumido em, ...). Confira o arquivo.")
    if "assumidoEm" not in m and "abertoEm" not in m:
        sys.exit('O arquivo precisa de "Assumido em" ou "Aberto em".')
    usadas = set(m.values())
    extras, descartadas = [], []
    for ci, h in enumerate(table[hi]):
        nome = "" if h is None else str(h).strip()
        if not nome or ci in usadas:
            continue
        (descartadas if sensivel(nome) else extras).append((nome, ci))
    if descartadas:
        print("Dados do cliente descartados:", ", ".join(n for n, _ in descartadas))
    if extras:
        print("Colunas extras mantidas:", ", ".join(n for n, _ in extras))
    out = []
    for line in table[hi + 1:]:
        o, any_ = {}, False
        for f in FIELDS:
            v = line[m[f]] if f in m and m[f] < len(line) else ""
            v = to_iso(v) if f in DATE_FIELDS else ("" if v is None else str(v).strip())
            if v == "-":
                v = ""
            o[f] = v
            any_ = any_ or bool(v)
        ex = {}
        for nome, ci in extras:
            v = line[ci] if ci < len(line) else ""
            v = to_iso(v) if isinstance(v, (datetime, date)) else ("" if v is None else str(v).strip())
            if v and v != "-":
                ex[nome] = v
        o["extras"] = ex
        if any_:
            out.append(o)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada", help="relatório exportado do OPA (.csv ou .xlsx)")
    ap.add_argument("-o", "--saida", default="data/atendimentos.csv")
    ap.add_argument("--acrescentar", action="store_true", help="junta ao arquivo de saída existente, sem duplicar")
    a = ap.parse_args()

    rows = canonicalize(read_table(a.entrada))
    key = lambda r: r["protocolo"] or r["id"]
    novas = len(rows)
    if a.acrescentar:
        try:
            old = canonicalize(read_table(a.saida))
        except (FileNotFoundError, SystemExit):
            old = []
        merged, extra = {}, []
        for r in old + rows:
            (merged.__setitem__(key(r), r) if key(r) else extra.append(r))
        rows = list(merged.values()) + extra
    with open(a.saida, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        extra_cols = sorted({k for r in rows for k in r.get("extras", {})})
        w.writerow([HEADERS[x] for x in FIELDS] + extra_cols)
        for r in rows:
            w.writerow([r[x] for x in FIELDS] + [r.get("extras", {}).get(k, "") for k in extra_cols])
    datas = sorted(r["assumidoEm"] or r["abertoEm"] for r in rows if (r["assumidoEm"] or r["abertoEm"]))
    print(f"OK: {len(rows)} linhas em {a.saida} ({novas} lidas do arquivo de entrada).")
    if datas:
        print(f"Período: {datas[0][:10]} a {datas[-1][:10]}")


if __name__ == "__main__":
    main()
